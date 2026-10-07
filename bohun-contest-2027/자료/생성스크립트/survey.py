import sys, random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

out = sys.argv[1]
fill_test = len(sys.argv) > 2 and sys.argv[2] == 'test'

FONT = '맑은 고딕'
NAVY = '1F3A5F'
f = lambda **k: Font(name=FONT, **k)
INPUT = PatternFill('solid', fgColor='FFF7CC')
HEAD = PatternFill('solid', fgColor=NAVY)
SUB = PatternFill('solid', fgColor='EEF2F7')
thin = Side(style='thin', color='B7C2D0')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

ITEMS = [
    '임시정부가 오늘의 대한민국과 어떻게 이어지는지 설명할 수 있다',
    '독립운동가·광복군의 희생이 내 삶과 관련 있다고 느낀다',
    '우리 동네 보훈 관련 장소를 알고 있다',
    '독립유공자와 가족에게 감사를 표현할 방법을 안다',
    '보훈 관련 활동에 참여할 의향이 있다',
    'AI·인터넷 역사 정보를 그대로 믿지 않고 확인한다',
    '역사 정보의 출처를 찾고 표기할 수 있다',
    '부정확한 역사 정보를 바로잡는 것도 보훈이다',
]
N = 40           # student rows
R0, R1 = 2, N + 1

wb = Workbook()

# ---------- 안내 ----------
g = wb.active; g.title = '안내'
g['A1'] = '사전·사후 설문 분석 시트 — 기억은 FACT로 지킨다'; g['A1'].font = f(bold=True, size=14, color=NAVY)
rows = [
    '사용 방법',
    '1. [사전] 시트와 [사후] 시트의 노란 칸에만 입력합니다. 학생번호(예: 30215 = 3학년 2반 15번)와 문항별 점수(1~5)를 넣습니다.',
    '2. 구글 설문 응답을 내려받아 붙여넣을 때도 열 순서를 학생번호, Q1~Q8로 맞춥니다. 학생 이름은 넣지 않습니다(개인정보 최소화).',
    '3. 사전과 사후는 학생번호로 자동 짝지어집니다. 행 순서가 달라도 됩니다. 한쪽만 응답한 학생은 대응 분석에서 빠집니다.',
    '4. [분석] 시트가 문항별 평균, 변화량, 대응표본 t검정(양측) 결과와 그래프를 자동으로 보여 줍니다.',
    '5. [대응] 시트는 계산용입니다. 수정하지 마세요.',
    '',
    '해석 기준',
    '· p < .05 이면 "통계적으로 유의한 변화"로 적을 수 있습니다. 표본(대응 n)이 10명 미만이면 결과를 참고용으로만 씁니다.',
    '· 공모안에는 [분석] 시트의 그래프 1장 + 영역 평균 2줄(보훈 의식 Q1~5, 정보 검증 태도 Q6~8)을 싣는 것을 권장합니다.',
    '',
    '입력 예시 (아래는 형식을 보여 주는 예시이며, 계산에 포함되지 않습니다)',
]
for i, txt in enumerate(rows, start=3):
    g.cell(i, 1, txt).font = f(bold=txt in ('사용 방법', '해석 기준') or txt.startswith('입력 예시'), size=11 if not txt.startswith(('사용', '해석', '입력')) else 12, color=NAVY if txt in ('사용 방법', '해석 기준') else '000000')
er = 3 + len(rows)
hdr = ['학생번호'] + [f'Q{i}' for i in range(1, 9)]
for j, v in enumerate(hdr, start=1):
    c = g.cell(er, j, v); c.font = f(bold=True, color='FFFFFF'); c.fill = HEAD; c.alignment = center; c.border = box
for j, v in enumerate([30215, 3, 2, 1, 2, 4, 2, 2, 3], start=1):
    c = g.cell(er + 1, j, v); c.font = f(); c.alignment = center; c.border = box; c.fill = INPUT
g.column_dimensions['A'].width = 12
for j in range(2, 10): g.column_dimensions[L(j)].width = 7
g.cell(er + 3, 1, '문항').font = f(bold=True, color=NAVY)
for i, q in enumerate(ITEMS, 1):
    g.cell(er + 3 + i, 1, f'Q{i}').font = f(bold=True)
    g.cell(er + 3 + i, 2, q).font = f()
g.cell(er + 12, 1, '영역: 보훈 의식 = Q1~Q5, 정보 검증 태도 = Q6~Q8').font = f(color='555555')

# ---------- 사전 / 사후 ----------
dv_score = None
for name in ('사전', '사후'):
    ws = wb.create_sheet(name)
    for j, v in enumerate(hdr, start=1):
        c = ws.cell(1, j, v); c.font = f(bold=True, color='FFFFFF'); c.fill = HEAD; c.alignment = center; c.border = box
    dv = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True,
                        showErrorMessage=True, errorTitle='점수 범위', error='1~5 사이 정수를 입력하세요.')
    ws.add_data_validation(dv)
    dv.add(f'B{R0}:I{R1}')
    for r in range(R0, R1 + 1):
        for j in range(1, 10):
            c = ws.cell(r, j); c.fill = INPUT; c.border = box; c.alignment = center; c.font = f()
    ws.column_dimensions['A'].width = 12
    for j in range(2, 10): ws.column_dimensions[L(j)].width = 7
    ws.freeze_panes = 'B2'
    if fill_test:
        random.seed(1 if name == '사전' else 2)
        ids = [30101 + k for k in range(25)]
        if name == '사후': ids = list(reversed(ids))[:23]   # different order + 2 missing
        for k, sid in enumerate(ids):
            ws.cell(R0 + k, 1, sid)
            for q in range(8):
                base = 2 if name == '사전' else 4
                ws.cell(R0 + k, 2 + q, max(1, min(5, base + random.choice([-1, 0, 0, 1]))))

# ---------- 대응 (paired differences) ----------
pw = wb.create_sheet('대응')
pw['A1'] = '학생번호'
for q in range(1, 9): pw.cell(1, 1 + q, f'Q{q} 변화')
pw.cell(1, 10, '보훈 의식 변화(Q1~5 평균)'); pw.cell(1, 11, '정보 검증 변화(Q6~8 평균)')
for j in range(1, 12):
    c = pw.cell(1, j); c.font = f(bold=True, color='FFFFFF'); c.fill = HEAD; c.alignment = center
for r in range(R0, R1 + 1):
    pw.cell(r, 1, f'=IF(사전!A{r}="","",사전!A{r})')
    for q in range(1, 9):
        col = L(1 + q)
        pre = f'사전!{col}{r}'
        post = f'INDEX(사후!{col}${R0}:{col}${R1},MATCH($A{r},사후!$A${R0}:$A${R1},0))'
        pw.cell(r, 1 + q, f'=IFERROR(IF(AND($A{r}<>"",ISNUMBER({pre}),ISNUMBER({post})),{post}-{pre},""),"")')
    pw.cell(r, 10, f'=IF(COUNT(B{r}:F{r})=5,AVERAGE(B{r}:F{r}),"")')
    pw.cell(r, 11, f'=IF(COUNT(G{r}:I{r})=3,AVERAGE(G{r}:I{r}),"")')
    for j in range(1, 12): pw.cell(r, j).font = f(color='555555'); pw.cell(r, j).alignment = center
pw.column_dimensions['A'].width = 12
for j in range(2, 12): pw.column_dimensions[L(j)].width = 11
pw.sheet_state = 'visible'

# ---------- 분석 ----------
a = wb.create_sheet('분석')
a['A1'] = '사전·사후 비교 분석'; a['A1'].font = f(bold=True, size=14, color=NAVY)
a['A2'] = '대응표본 t검정(양측). 대응 n = 사전·사후 모두 응답한 학생 수. 노란 칸 입력은 [사전]·[사후] 시트에서 합니다.'; a['A2'].font = f(color='555555')
cols = ['문항', '내용', '사전 평균', '사후 평균', '평균 변화(대응)', '대응 n', '변화의 표준편차', 't', 'p (양측)', '판정']
HR = 4
for j, v in enumerate(cols, start=1):
    c = a.cell(HR, j, v); c.font = f(bold=True, color='FFFFFF'); c.fill = HEAD; c.alignment = center; c.border = box

def row(r, label, desc, pre_rng, post_rng, diff_rng, fill=None):
    a.cell(r, 1, label); a.cell(r, 2, desc)
    a.cell(r, 3, f'=IFERROR(AVERAGE({pre_rng}),"-")')
    a.cell(r, 4, f'=IFERROR(AVERAGE({post_rng}),"-")')
    a.cell(r, 5, f'=IFERROR(AVERAGE({diff_rng}),"-")')
    a.cell(r, 6, f'=COUNT({diff_rng})')
    a.cell(r, 7, f'=IF(F{r}>=2,STDEV({diff_rng}),"-")')
    a.cell(r, 8, f'=IFERROR(E{r}/(G{r}/SQRT(F{r})),"-")')
    a.cell(r, 9, f'=IFERROR(TDIST(ABS(H{r}),F{r}-1,2),"-")')
    a.cell(r, 10, f'=IF(ISNUMBER(I{r}),IF(I{r}<0.05,"유의함 (p<.05)","유의하지 않음"),"자료 부족")')
    for j in range(1, 11):
        c = a.cell(r, j); c.border = box; c.font = f(bold=fill is not None); c.alignment = center if j != 2 else Alignment(vertical='center', wrap_text=True)
        if fill: c.fill = fill
    for j in (3, 4, 5, 7, 8): a.cell(r, j).number_format = '0.00'
    a.cell(r, 9).number_format = '0.000'

for q in range(1, 9):
    col = L(1 + q)
    row(HR + q, f'Q{q}', ITEMS[q - 1], f'사전!{col}{R0}:{col}{R1}', f'사후!{col}{R0}:{col}{R1}', f'대응!{col}{R0}:{col}{R1}')
r = HR + 9
row(r, '영역', '보훈 의식 (Q1~Q5)', f'사전!B{R0}:F{R1}', f'사후!B{R0}:F{R1}', f'대응!J{R0}:J{R1}', fill=SUB)
row(r + 1, '영역', '정보 검증 태도 (Q6~Q8)', f'사전!G{R0}:I{R1}', f'사후!G{R0}:I{R1}', f'대응!K{R0}:K{R1}', fill=SUB)
a.cell(r + 2, 1, '※ 영역 행의 사전·사후 평균은 해당 문항 전체 응답의 평균, 평균 변화와 t검정은 학생별 영역 평균의 변화로 계산합니다.').font = f(color='555555', size=9)
a.column_dimensions['A'].width = 7
a.column_dimensions['B'].width = 46
for j, w in zip('CDEFGHIJ', [10, 10, 13, 8, 13, 8, 10, 15]): a.column_dimensions[j].width = w

ch = BarChart(); ch.type = 'bar'; ch.grouping = 'clustered'
ch.title = '문항별 사전·사후 평균 (5점 척도)'
ch.y_axis.title = '평균 점수'; ch.y_axis.scaling.min = 1; ch.y_axis.scaling.max = 5; ch.y_axis.majorUnit = 1
ch.y_axis.majorGridlines = None
ch.x_axis.scaling.orientation = 'maxMin'
data = Reference(a, min_col=3, max_col=4, min_row=HR, max_row=HR + 8)
cats = Reference(a, min_col=1, min_row=HR + 1, max_row=HR + 8)
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
for s, color in zip(ch.series, ['2A78D6', 'EB6834']):
    s.graphicalProperties.solidFill = color
    s.graphicalProperties.line.solidFill = color
ch.gapWidth = 80
ch.legend.position = 't'
ch.height = 10; ch.width = 18
a.add_chart(ch, f'B{r + 4}')
a.page_setup.orientation = 'landscape'
a.sheet_properties.pageSetUpPr.fitToPage = True
a.page_setup.fitToWidth = 1
a.page_setup.fitToHeight = 0
g.page_setup.orientation = 'landscape'
g.sheet_properties.pageSetUpPr.fitToPage = True
g.page_setup.fitToWidth = 1
g.page_setup.fitToHeight = 0

wb.move_sheet('분석', offset=-3)  # order: 안내, 분석, 사전, 사후, 대응
wb.save(out)
print('saved', out)
